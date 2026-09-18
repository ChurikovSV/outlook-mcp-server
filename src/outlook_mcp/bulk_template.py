from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .markdown_draft import create_draft_from_markdown
from .models import EmailRequest


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    return _stringify(value)


def _load_rows_from_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(handle, dialect=dialect)
        if not reader.fieldnames:
            raise ValueError("CSV file has no header row")
        headers = [_normalize_header(name) for name in reader.fieldnames]
        if any(not name for name in headers):
            raise ValueError("CSV contains an empty column name")
        rows: list[dict[str, str]] = []
        for raw in reader:
            row = {header: _stringify(raw.get(original)) for header, original in zip(headers, reader.fieldnames)}
            if any(row.values()):
                rows.append(row)
        return rows


def _load_rows_from_xlsx(path: Path, sheet: str | None) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise ValueError(
                    f"Sheet '{sheet}' not found. Available sheets: {', '.join(workbook.sheetnames)}"
                )
            worksheet = workbook[sheet]
        else:
            worksheet = workbook[workbook.sheetnames[0]]

        rows_iter = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as exc:
            raise ValueError("XLSX file is empty") from exc

        headers = [_normalize_header(value) for value in header_row]
        if any(not name for name in headers):
            raise ValueError("XLSX contains an empty column name")
        if len(set(name.casefold() for name in headers)) != len(headers):
            raise ValueError("XLSX contains duplicate column names")

        rows: list[dict[str, str]] = []
        for values in rows_iter:
            row = {
                header: _stringify(values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
            }
            if any(row.values()):
                rows.append(row)
        return rows
    finally:
        workbook.close()


def _load_rows(file_path: str, sheet: str | None) -> list[dict[str, str]]:
    path = Path(file_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Recipient data file not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _load_rows_from_csv(path)
    if suffix == ".xlsx":
        return _load_rows_from_xlsx(path, sheet)
    raise ValueError("Recipient data file must be .csv or .xlsx")


def _find_column(row: dict[str, str], requested: str, aliases: tuple[str, ...] = ()) -> str | None:
    by_casefold = {key.casefold(): key for key in row}
    for candidate in (requested, *aliases):
        actual = by_casefold.get(candidate.casefold())
        if actual:
            return actual
    return None


def _split_addresses(value: str) -> list[str]:
    if not value.strip():
        return []
    normalized = value.replace(",", ";")
    return [part.strip() for part in normalized.split(";") if part.strip()]


def create_bulk_drafts_from_template(
    recipient_file: str,
    template_name: str,
    subject: str = "",
    email_column: str = "email",
    subject_column: str = "subject",
    sheet: str | None = None,
    from_email: str | None = None,
    attachments: list[str] | None = None,
    uploaded_attachments: list[dict] | None = None,
) -> dict:
    """
    Create one personalized Outlook draft per data row.

    Every column is exposed to the Markdown template as {{column_name}}.
    Reserved columns may additionally control delivery fields:
    email, subject/Тема, cc, bcc, from_email.
    """
    rows = _load_rows(recipient_file, sheet)
    if not rows:
        raise ValueError("Recipient data file has no data rows")

    created = 0
    failed: list[dict[str, object]] = []
    drafts: list[dict[str, object]] = []

    for index, row in enumerate(rows, start=2):
        try:
            email_key = _find_column(row, email_column)
            if not email_key or not row[email_key]:
                raise ValueError(f"Missing recipient email in column '{email_column}'")

            subject_key = _find_column(row, subject_column, aliases=("Тема",))
            row_subject = row.get(subject_key, "") if subject_key else ""
            effective_subject = row_subject or subject
            if not effective_subject:
                raise ValueError(
                    f"Missing subject: provide '{subject_column}'/'Тема' column or subject argument"
                )

            cc_key = _find_column(row, "cc")
            bcc_key = _find_column(row, "bcc")
            sender_key = _find_column(row, "from_email")

            effective_sender = row.get(sender_key, "") if sender_key else ""
            effective_sender = effective_sender or from_email

            variables = dict(row)
            # Canonical aliases are always available even when localized/custom service
            # column names were used.
            variables.setdefault("email", row[email_key])
            variables.setdefault("subject", effective_subject)
            if subject_key:
                variables.setdefault("Тема", row.get(subject_key, ""))
            variables.setdefault("cc", row.get(cc_key, "") if cc_key else "")
            variables.setdefault("bcc", row.get(bcc_key, "") if bcc_key else "")
            variables.setdefault("from_email", effective_sender or "")

            request = EmailRequest(
                to=[row[email_key]],
                cc=_split_addresses(row.get(cc_key, "") if cc_key else ""),
                bcc=_split_addresses(row.get(bcc_key, "") if bcc_key else ""),
                from_email=effective_sender,
                subject=effective_subject,
                attachments=attachments or [],
                uploaded_attachments=uploaded_attachments or [],
            )
            result = create_draft_from_markdown(
                request=request,
                template_name=template_name,
                variables=variables,
            )
            created += 1
            drafts.append(
                {
                    "row": index,
                    "email": row[email_key],
                    "subject": effective_subject,
                    "entry_id": result.get("entry_id"),
                    "from_email": result.get("from_email"),
                }
            )
        except Exception as exc:
            failed.append(
                {
                    "row": index,
                    "email": row.get(_find_column(row, email_column) or "", ""),
                    "error": str(exc),
                }
            )

    return {
        "status": "completed" if not failed else "completed_with_errors",
        "template_name": template_name,
        "source_file": str(Path(recipient_file).expanduser().resolve()),
        "sheet": sheet,
        "total": len(rows),
        "created": created,
        "failed": failed,
        "drafts": drafts,
        "column_mapping": "every_column_is_template_variable",
        "reserved_columns": ["email", "subject", "Тема", "cc", "bcc", "from_email"],
    }
